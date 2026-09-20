import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from report_generator.styles import setup_document, PRIMARY_COLOR, DARK_TEXT
from report_generator.front_matter import (
    add_cover_page, add_certificate, add_declaration,
    add_company_certificate_placeholder, add_company_profile,
    add_acknowledgements, add_abstract, add_table_of_contents,
    add_list_of_figures, add_list_of_tables
)
from report_generator.chapter1_3 import add_chapter1, add_chapter2, add_chapter3
from report_generator.chapter4 import add_chapter4
from report_generator.chapter5 import add_chapter5
from report_generator.chapter6 import add_chapter6
from report_generator.chapter7_12 import (
    add_chapter7, add_chapter8, add_chapter9, add_chapter11, add_chapter12
)

def setup_footer(section, is_roman=False):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    r_prefix = p.add_run("MediMind AI | Page ")
    r_prefix.font.name = "Times New Roman"
    r_prefix.font.size = Pt(9.5)
    r_prefix.font.italic = True
    r_prefix.font.color.rgb = RGBColor(100, 116, 139)
    
    r_pg = p.add_run()
    r_pg.font.name = "Times New Roman"
    r_pg.font.size = Pt(9.5)
    r_pg.font.italic = True
    r_pg.font.color.rgb = RGBColor(100, 116, 139)
    fldSimple = parse_xml(f"<w:fldSimple {nsdecls('w')} w:instr='PAGE'/>")
    r_pg._r.append(fldSimple)

def generate_report():
    print("Initializing document...")
    doc = setup_document()

    # Section 1: Front Matter
    sec1 = doc.sections[0]
    sec1.top_margin = Inches(1.0)
    sec1.bottom_margin = Inches(1.0)
    sec1.left_margin = Inches(1.5)
    sec1.right_margin = Inches(1.0)
    
    pnt1 = OxmlElement("w:pgNumType")
    pnt1.set(qn("w:fmt"), "lowerRoman")
    sec1._sectPr.append(pnt1)

    print("Adding Cover Page...")
    add_cover_page(doc)

    print("Adding Certificate...")
    add_certificate(doc)

    print("Adding Declaration...")
    add_declaration(doc)

    print("Adding Company Internship Certificate Placeholder...")
    add_company_certificate_placeholder(doc)

    print("Adding Company Profile & External Guide Details...")
    add_company_profile(doc)

    print("Adding Acknowledgements...")
    add_acknowledgements(doc)

    print("Adding Abstract & Keywords...")
    add_abstract(doc)

    print("Adding Table of Contents...")
    add_table_of_contents(doc)

    print("Adding List of Figures...")
    add_list_of_figures(doc)

    print("Adding List of Tables...")
    add_list_of_tables(doc)

    # Section 2: Main Chapters
    print("Configuring Body Section...")
    sec2 = doc.add_section()
    sec2.top_margin = Inches(1.0)
    sec2.bottom_margin = Inches(1.0)
    sec2.left_margin = Inches(1.5)
    sec2.right_margin = Inches(1.0)
    
    pnt2 = OxmlElement("w:pgNumType")
    pnt2.set(qn("w:fmt"), "decimal")
    pnt2.set(qn("w:start"), "1")
    sec2._sectPr.append(pnt2)
    
    setup_footer(sec2, is_roman=False)

    print("Adding Chapter 1 (Introduction)...")
    add_chapter1(doc)

    print("Adding Chapter 2 (Literature Survey / Existing System)...")
    add_chapter2(doc)

    print("Adding Chapter 3 (Software Requirement Analysis)...")
    add_chapter3(doc)

    print("Adding Chapter 4 (Software Design & UML/DFD/ER Diagrams)...")
    add_chapter4(doc)

    print("Adding Chapter 5 (Technology / Methodology & Modules)...")
    add_chapter5(doc)

    print("Adding Chapter 6 (Coding / Implementation & Real Source Listings)...")
    add_chapter6(doc)

    print("Adding Chapter 7 (Testing & Test Cases)...")
    add_chapter7(doc)

    print("Adding Chapter 8 (Output Screens & Results)...")
    add_chapter8(doc)

    print("Adding Chapter 9 (Conclusion, Limitations & Future Work)...")
    add_chapter9(doc)

    print("Adding Chapter 11 (References)...")
    add_chapter11(doc)

    print("Adding Chapter 12 (Appendices)...")
    add_chapter12(doc)

    output_filename = "MediMind_AI_Internship_Report.docx"
    doc.save(output_filename)
    print(f"Report successfully saved to: {output_filename}")

    ref_copy = os.path.join("report_reference", output_filename)
    shutil.copy(output_filename, ref_copy)
    print(f"Backup copy saved to: {ref_copy}")

    size_kb = os.path.getsize(output_filename) / 1024
    print(f"Final Report File Size: {size_kb:.2f} KB")

if __name__ == "__main__":
    generate_report()