import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from report_generator.styles import (
    add_heading1, add_heading2, add_heading3, add_paragraph,
    add_bullet, add_table
)

def add_chapter1(doc):
    add_heading1(doc, "CHAPTER 1\nINTRODUCTION", centered=True)

    add_heading2(doc, "1.1 Introduction to Healthcare Informatics & MediMind AI")
    add_paragraph(
        doc,
        "Healthcare informatics bridges information technology and healthcare services to improve health literacy, "
        "facilitate early symptom awareness, and assist users in understanding medical information. In everyday life, "
        "individuals frequently encounter diagnostic reports and health documentation that contain complex medical terms, "
        "standard reference ranges, and laboratory markers that are difficult to understand without medical training."
    )
    add_paragraph(
        doc,
        "MediMind AI is an AI-powered healthcare assistant application developed during an internship at SkillDzire. "
        "The project integrates an asynchronous backend built with FastAPI and Python, a modern Single Page Application (SPA) "
        "frontend developed with React and Tailwind CSS, and cloud-based Large Language Model (LLM) inference through the Groq Cloud API "
        "using the 'openai/gpt-oss-20b' model."
    )
    add_paragraph(
        doc,
        "A central feature of MediMind AI is its in-memory medical report PDF analysis capability. Utilizing PyMuPDF (fitz), "
        "the application processes uploaded PDF reports directly from memory byte streams without writing temporary files to disk, "
        "preventing Windows file-lock issues and ensuring privacy. Furthermore, MediMind AI incorporates bilingual support in English "
        "and Telugu, complete with a virtual on-screen Telugu keyboard to facilitate easy interaction for regional language users."
    )

    add_heading2(doc, "1.2 Problem Statement")
    add_paragraph(
        doc,
        "Users frequently encounter significant barriers when seeking clear health information and reviewing laboratory diagnostic reports:"
    )
    add_bullet(
        doc,
        "Medical laboratory reports contain complex clinical terms and reference ranges that are confusing for non-medical individuals, "
        "often causing unnecessary anxiety.",
        "Complexity of Diagnostic Reports: "
    )
    add_bullet(
        doc,
        "Most digital health tools and informational portals are offered exclusively in English, creating accessibility hurdles for users "
        "who prefer regional languages such as Telugu.",
        "Language Accessibility Barriers: "
    )
    add_bullet(
        doc,
        "General internet searches often produce unverified, contradictory, or alarming health information without clear disclaimers or triage guidance.",
        "Lack of Structured Health Explanations: "
    )
    add_bullet(
        doc,
        "Traditional web file-upload workflows often write uploaded files to disk storage, creating potential file cleanup issues and OS file concurrency locks.",
        "Temporary File Management: "
    )

    add_heading2(doc, "1.3 Objectives of the Project")
    add_paragraph(
        doc,
        "The primary objectives of the MediMind AI project are:"
    )
    add_bullet(doc, "Develop an asynchronous RESTful API backend using FastAPI, providing endpoints for authentication, AI querying, PDF report analysis, and activity history.", "FastAPI Backend Architecture: ")
    add_bullet(doc, "Implement direct in-memory text extraction from uploaded PDF reports using PyMuPDF (fitz), eliminating temporary file creation on disk.", "In-Memory PDF Extraction: ")
    add_bullet(doc, "Integrate the Groq Cloud API using the 'openai/gpt-oss-20b' model with carefully structured system prompts and safety guardrails.", "Cloud LLM Integration & Safety Rules: ")
    add_bullet(doc, "Provide bilingual user support (English and Telugu) with client-side localization and an interactive virtual on-screen Telugu keyboard.", "Bilingual Support & Telugu Keyboard: ")
    add_bullet(doc, "Implement token-based authentication using JSON Web Tokens (JWT) and secure password hashing using bcrypt via Passlib.", "Secure User Authentication: ")
    add_bullet(doc, "Store user consultation activities in an SQLite database using SQLAlchemy ORM to provide user dashboard statistics and recent activity history.", "Activity Tracking & Analytics: ")

    add_heading2(doc, "1.4 Scope of the Project")
    add_paragraph(
        doc,
        "The operational and clinical scope of MediMind AI is defined as follows:"
    )
    add_bullet(doc, "Provides AI-based symptom checks, medical report PDF analysis, medicine information lookup, and general health conversational assistance.", "Functional Scope: ")
    add_bullet(doc, "MediMind AI provides educational health information only. It explicitly does not diagnose medical conditions, does not prescribe medications, does not recommend dosages, and does not replace consultation with a licensed healthcare professional.", "Clinical Safety Boundaries: ")
    add_bullet(doc, "Responsive web application accessible via standard desktop and mobile web browsers.", "Platform Scope: ")

    doc.add_page_break()

def add_chapter2(doc):
    add_heading1(doc, "CHAPTER 2\nLITERATURE SURVEY / EXISTING SYSTEM", centered=True)

    add_heading2(doc, "2.1 Existing System Overview")
    add_paragraph(
        doc,
        "In existing healthcare information solutions, users typically rely on online medical encyclopedias (e.g., WebMD, Healthline), "
        "telemedicine platforms (e.g., Practo, 1mg), or general-purpose artificial intelligence conversational tools. "
        "These platforms provide valuable clinical libraries and appointment scheduling, but operate with distinct characteristics:"
    )
    add_bullet(doc, "Websites provide static articles and rule-based symptom trees requiring users to self-select predefined checkboxes.", "Medical Encyclopedias: ")
    add_bullet(doc, "Focus on doctor appointments and medicine ordering; diagnostic reports are uploaded as passive attachments for human physician review without automated explanation.", "Telemedicine Portals: ")
    add_bullet(doc, "Accept free-text prompts, but lack specialized, consistent clinical output templates or integrated in-memory PDF extraction pipelines tailored for patient report reading.", "General-Purpose AI Chatbots: ")

    add_heading2(doc, "2.2 Limitations of Existing Systems")
    add_paragraph(
        doc,
        "Key limitations identified in current consumer health tools include:"
    )
    add_bullet(doc, "Users must manually interpret tabular values and reference intervals, often misinterpreting benign variations.", "Lack of Automated Report Interpretation: ")
    add_bullet(doc, "Most platforms offer interfaces primarily in English, with minimal support for interactive vernacular typing in regional scripts like Telugu.", "Regional Language Barriers: ")
    add_bullet(doc, "Many tools do not enforce rigid structured output formatting, occasionally generating overly broad or alarmist information.", "Unstructured Responses: ")
    add_bullet(doc, "Many document-handling backends save uploaded files to local disk storage, which can create file-lock conflicts and storage management overhead.", "Server-Side File Persistence Overhead: ")

    add_heading2(doc, "2.3 Literature Survey")
    add_paragraph(
        doc,
        "Research in clinical Natural Language Processing (NLP) and healthcare user experience has explored various dimensions "
        "of AI-assisted health communication:"
    )
    add_paragraph(
        doc,
        "Scholarly investigations by researchers in biomedical informatics have studied the performance of large language models "
        "on standardized medical question-answering benchmarks. These studies emphasize that while foundation models contain broad biomedical "
        "knowledge, applying structured prompt engineering and explicit negative constraints (such as disclaimers and refusal to prescribe) "
        "is critical to maintaining safety.",
        bold_prefix="1. Clinical NLP & Safety Constraints: "
    )
    add_paragraph(
        doc,
        "Studies on digital health accessibility in regional and developing demographics emphasize that translating medical information into "
        "local languages requires maintaining clarity. Mechanical translations often distort clinical meanings; retaining familiar medical terms "
        "in brackets alongside simplified native explanations significantly improves patient understanding.",
        bold_prefix="2. Vernacular Healthcare Communication: "
    )
    add_paragraph(
        doc,
        "Software engineering literature on Python document processing indicates that processing digital documents directly from memory byte streams "
        "(via libraries like PyMuPDF) delivers higher throughput and avoids temporary file I/O compared to older disk-dependent parsers.",
        bold_prefix="3. Document Parsing & Stream Handling: "
    )
    add_paragraph(
        doc,
        "It is important to clearly distinguish external academic literature from the actual implementation of MediMind AI. "
        "MediMind AI does not implement proprietary clinical models (such as Med-PaLM) or custom-trained biomedical neural networks. "
        "Instead, MediMind AI is an applied, production-oriented healthcare assistant leveraging the Groq API with the 'openai/gpt-oss-20b' model, "
        "combined with in-memory PyMuPDF stream parsing, structured clinical prompt guardrails, and bilingual English-Telugu localization.",
        bold_prefix="Distinction from Research Models: "
    )

    add_heading2(doc, "2.4 Proposed System & Key Advantages")
    add_paragraph(
        doc,
        "MediMind AI addresses the aforementioned limitations through a targeted, privacy-conscious full-stack design:"
    )
    add_bullet(doc, "Extracts readable text directly from PDF bytes in memory, completely bypassing disk writes and eliminating file-lock concurrency issues on Windows.", "In-Memory PDF Stream Processing: ")
    add_bullet(doc, "Leverages the Groq Cloud API running 'openai/gpt-oss-20b', generating comprehensive, safe health explanations with low response latency.", "Cloud LLM Integration: ")
    add_bullet(doc, "System prompts enforce structured sections (Summary, Findings, Abnormal Values, Health Suggestions, Doctor Consultation Advice, Disclaimers) and prohibit dosage prescriptions.", "Structured Safety Rules: ")
    add_bullet(doc, "Provides both interface localization and a dedicated on-screen Telugu keyboard for effortless vernacular input.", "Telugu Language Support & Virtual Keyboard: ")
    add_bullet(doc, "Uses bcrypt password hashing and JWT access tokens to ensure private, isolated user consultation logs.", "JWT-Secured Sessions & Activity Logging: ")

    doc.add_page_break()

def add_chapter3(doc):
    add_heading1(doc, "CHAPTER 3\nSOFTWARE REQUIREMENT ANALYSIS", centered=True)

    add_heading2(doc, "3.1 Software Requirements")
    add_paragraph(
        doc,
        "Table 3.1 lists the verified software components, frameworks, and library versions utilized in the MediMind AI codebase."
    )

    headers = ["Layer / Subsystem", "Technology / Framework", "Version", "Role in Project"]
    rows = [
        ["Operating System", "Microsoft Windows", "10 / 11 (64-bit)", "Development and execution environment"],
        ["Backend Runtime", "Python", "3.14.x", "Backend programming language runtime"],
        ["Backend Web Framework", "FastAPI", "0.141.0", "Asynchronous REST API framework"],
        ["ASGI Web Server", "Uvicorn", "0.52.0", "ASGI server for running FastAPI"],
        ["Object Relational Mapper", "SQLAlchemy", "2.0.52", "ORM for database modeling and query execution"],
        ["Relational Database", "SQLite 3", "3.x", "File-based relational database (medimind.db)"],
        ["PDF Extraction Library", "PyMuPDF (fitz)", "1.28.0", "In-memory extraction of text from PDF bytes"],
        ["AI / LLM Cloud SDK", "Groq Python SDK", "1.6.0", "Client library for calling Groq Cloud LLM API"],
        ["Authentication / JWT", "python-jose", "3.5.0", "JWT encoding, decoding, and signature verification"],
        ["Password Cryptography", "Passlib & Bcrypt", "1.7.4 / 4.0.1", "Bcrypt-based secure password hashing"],
        ["Environment Configuration", "python-dotenv", "1.2.2", "Loading environment variables from .env file"],
        ["Frontend Library", "React.js", "19.x", "Component-based UI library for frontend SPA"],
        ["Frontend Build Tool", "Vite", "6.x", "Development server and frontend asset bundler"],
        ["Styling Framework", "Tailwind CSS", "3.4.x", "Utility-first CSS styling framework"],
        ["HTTP Client", "Axios", "1.7.x", "Promise-based HTTP client with request interceptors"],
        ["Icons Library", "Lucide React", "0.4.x", "Vector iconography for user interface elements"],
        ["Notifications", "React Toastify", "10.x", "Non-blocking toast alerts for user notifications"]
    ]
    add_table(doc, headers, rows, [1.4, 1.6, 0.9, 2.8])
    add_paragraph(doc, "Table 3.1: Software Technology Stack & Library Versions", italic=True)

    add_heading2(doc, "3.2 Hardware Requirements")
    add_paragraph(
        doc,
        "Table 3.2 defines the general hardware requirements for running and testing the application."
    )

    h_headers = ["Hardware Component", "Minimum Development Specification", "Client Device Specification"]
    h_rows = [
        ["Processor (CPU)", "Intel Core i3 / AMD Ryzen 3 or equivalent", "Modern dual-core processor or smartphone CPU"],
        ["System Memory (RAM)", "8 GB RAM", "2 GB - 4 GB RAM (Desktop / Mobile)"],
        ["Storage (Disk)", "500 MB free disk space for project and dependencies", "Standard browser cache"],
        ["Network Connection", "Internet connection required for Groq Cloud API calls", "Broadband or cellular mobile internet (4G/5G)"],
        ["Display Resolution", "1280 x 720 minimum screen resolution", "Responsive across mobile, tablet, and desktop screens"]
    ]
    add_table(doc, h_headers, h_rows, [1.8, 2.4, 2.5])
    add_paragraph(doc, "Table 3.2: Hardware System Requirements", italic=True)

    add_heading2(doc, "3.3 Functional Requirements Definitions")
    add_paragraph(
        doc,
        "The functional requirements reflect the actual implemented capabilities in the MediMind AI codebase:"
    )
    add_bullet(doc, "Users can register with name, email, and password (minimum 6 characters). Users can log in with credentials to receive a signed JWT access token (60-minute expiry). Users can reset passwords via the forgot-password endpoint.", "FR-1: User Registration, Authentication & Password Reset: ")
    add_bullet(doc, "Authenticated users can submit health-related questions and receive concise, educational AI answers formatted in markdown.", "FR-2: AI Health Chat: ")
    add_bullet(doc, "Users can submit natural language descriptions of symptoms and receive structured assessments covering Possible Common Causes, General Self-Care, When to Consult a Doctor, and Emergency Warning Signs.", "FR-3: AI Symptom Checker: ")
    add_bullet(doc, "Users can upload medical laboratory reports in PDF format. The system extracts text directly from memory using PyMuPDF and returns structured clinical sections (Summary, Findings, Abnormal Values, Health Suggestions, Doctor Consultation Advice, Disclaimers).", "FR-4: Medical Report PDF Analyzer: ")
    add_bullet(doc, "Users can search medicine names to receive structured educational information regarding Main Uses, Common Side Effects, Precautions, and Safety Advice without dosage prescriptions.", "FR-5: Medicine Information Lookup: ")
    add_bullet(doc, "Every user consultation is recorded in the SQLite database. Users can view their past activity list and inspect real-time dashboard counts (Total Activities, Reports Analyzed, Symptom Checks, AI Chats, Medicine Searches).", "FR-6: Activity History & Dashboard Analytics: ")
    add_bullet(doc, "Users can switch the interface language between English and Telugu at any time. Users can type in Telugu using an interactive virtual on-screen Telugu keyboard.", "FR-7: Multilingual Support & Virtual Telugu Keyboard: ")

    add_heading2(doc, "3.4 Non-Functional Requirements Definitions")
    add_paragraph(
        doc,
        "The non-functional requirements governing the application are:"
    )
    add_bullet(doc, "Passwords must be securely hashed using bcrypt before database storage. Protected routes must require valid Bearer JWT tokens in request headers.", "NFR-1: Security & Credential Protection: ")
    add_bullet(doc, "The application interface must be responsive across desktop, tablet, and mobile displays, with accessible visual contrast and toast notifications.", "NFR-2: Usability & User Experience: ")
    add_bullet(doc, "The backend must handle invalid file uploads (non-PDF files, empty files, unreadable files) gracefully and return descriptive HTTP 400 error messages.", "NFR-3: Error Handling & Robustness: ")
    add_bullet(doc, "The AI inference must follow strict safety guardrails: no drug prescriptions, no dosage instructions, no definitive diagnosis, and mandatory inclusion of educational disclaimers.", "NFR-4: Clinical Safety & Ethical Guardrails: ")

    doc.add_page_break()