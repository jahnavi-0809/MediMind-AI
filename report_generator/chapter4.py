import os
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from report_generator.styles import (
    add_heading1, add_heading2, add_heading3, add_paragraph,
    add_bullet, add_table, add_figure
)

DIAGRAMS_DIR = "report_diagrams"

def add_chapter4(doc):
    add_heading1(doc, "CHAPTER 4\nSOFTWARE DESIGN", centered=True)

    add_heading2(doc, "4.1 System Architecture")
    add_paragraph(
        doc,
        "The architecture of MediMind AI follows a clean three-tier client-server model, separating user presentation, "
        "application business logic, and data storage. Figure 4.1 illustrates the architectural blueprint of the application."
    )

    fig_arch = os.path.join(DIAGRAMS_DIR, "fig4_1_system_architecture.png")
    if os.path.exists(fig_arch):
        add_figure(doc, fig_arch, "Figure 4.1: MediMind AI - Three-Tier System Architecture")

    add_paragraph(
        doc,
        "1. Presentation Tier (Frontend Client): Developed as a Single Page Application (SPA) using React 19, "
        "Vite, and Tailwind CSS. It handles page routing via React Router DOM, maintains shared state with React Context API "
        "(AuthContext and LanguageContext), renders an on-screen virtual Telugu keyboard component, and communicates with the backend "
        "using Axios with request interceptors for JWT Bearer tokens.",
        bold_prefix="• "
    )
    add_paragraph(
        doc,
        "2. Application Tier (REST API Backend): Built with Python and FastAPI, running on the Uvicorn ASGI server. "
        "The backend organizes functionality into modular APIRouters (/auth, /ai, /history). Business logic is handled through "
        "specialized services: auth_service.py for password hashing and JWT issuance, pdf_service.py for in-memory PyMuPDF extraction, "
        "and llm_service.py for Groq API integration and prompt construction.",
        bold_prefix="• "
    )
    add_paragraph(
        doc,
        "3. Data & AI Inference Tier: Consists of an embedded SQLite relational database (medimind.db) managed via SQLAlchemy ORM "
        "for user credentials and activity logging, and the Groq Cloud LLM API running the 'openai/gpt-oss-20b' model for generating health explanations.",
        bold_prefix="• "
    )

    add_heading2(doc, "4.2 Data Flow Diagrams (DFDs)")
    add_paragraph(
        doc,
        "Data Flow Diagrams describe the logical flow of data through the MediMind AI application at multiple levels of detail."
    )

    add_heading3(doc, "4.2.1 DFD Level 0 - Context Diagram")
    add_paragraph(
        doc,
        "The Level 0 Context Diagram represents the high-level boundary of the MediMind AI system. The primary user interacts with the system "
        "to submit authentication credentials, symptom descriptions, PDF reports, or medicine queries. The system processes requests using the "
        "Groq Cloud API and persists user data and activity records in the SQLite database. Figure 4.2 illustrates this flow."
    )

    fig_dfd0 = os.path.join(DIAGRAMS_DIR, "fig4_2_dfd_level_0.png")
    if os.path.exists(fig_dfd0):
        add_figure(doc, fig_dfd0, "Figure 4.2: Data Flow Diagram (DFD) Level 0 - Context Diagram")

    add_heading3(doc, "4.2.2 DFD Level 1 - System Flow Diagram")
    add_paragraph(
        doc,
        "The Level 1 DFD decomposes the system into five primary functional processes: 1.0 User Authentication & Token Issuance, "
        "2.0 Symptom Checker & Triage Analysis, 3.0 Medical Report PDF Analyzer, 4.0 Medicine Information Retrieval, and 5.0 Activity History "
        "& Stats Engine. It shows interactions with data stores D1 (Users) and D2 (Activities) and the Groq AI API. Figure 4.3 shows the details."
    )

    fig_dfd1 = os.path.join(DIAGRAMS_DIR, "fig4_3_dfd_level_1.png")
    if os.path.exists(fig_dfd1):
        add_figure(doc, fig_dfd1, "Figure 4.3: Data Flow Diagram (DFD) Level 1 - System Flow Diagram")

    add_heading3(doc, "4.2.3 DFD Level 2 - Medical Report Analyzer (Process 3.0)")
    add_paragraph(
        doc,
        "The Level 2 DFD outlines the step-by-step processing of diagnostic reports in Process 3.0: 3.1 PDF Upload & Header Validation, "
        "3.2 PyMuPDF In-Memory Stream Text Extraction, 3.3 Multilingual Prompt Construction, 3.4 Groq LLM Inference, 3.5 Activity Logging to D2, "
        "and 3.6 Structured UI Rendering. Figure 4.4 illustrates this sequence."
    )

    fig_dfd2 = os.path.join(DIAGRAMS_DIR, "fig4_4_dfd_level_2.png")
    if os.path.exists(fig_dfd2):
        add_figure(doc, fig_dfd2, "Figure 4.4: Data Flow Diagram (DFD) Level 2 - Medical Report Analyzer (Process 3.0)")

    add_heading2(doc, "4.3 Unified Modeling Language (UML) Diagrams")
    add_paragraph(
        doc,
        "UML diagrams provide standard structural and behavioral representations of the system components and interactions."
    )

    add_heading3(doc, "4.3.1 UML Use Case Diagram")
    add_paragraph(
        doc,
        "Figure 4.5 illustrates the Use Case Diagram for MediMind AI, showing the user's interaction with the main system features: "
        "Registration, Login, AI Health Chat, Symptom Checking, PDF Report Analysis, Medicine Information Lookup, History & Stats, and Language Switching."
    )

    fig_uc = os.path.join(DIAGRAMS_DIR, "fig4_5_use_case.png")
    if os.path.exists(fig_uc):
        add_figure(doc, fig_uc, "Figure 4.5: UML Use Case Diagram - MediMind AI Healthcare Platform")

    add_heading3(doc, "4.3.2 UML Class Diagram")
    add_paragraph(
        doc,
        "The Class Diagram in Figure 4.6 details the backend data models (User, Activity) and service modules (AuthService, LLMService, PDFService, "
        "Database). The User entity maintains a 1-to-Many association with the Activity entity."
    )

    fig_cls = os.path.join(DIAGRAMS_DIR, "fig4_6_class_diagram.png")
    if os.path.exists(fig_cls):
        add_figure(doc, fig_cls, "Figure 4.6: UML Class Diagram - Backend Data Models & Core Services")

    add_heading3(doc, "4.3.3 UML Interaction Diagrams")
    add_paragraph(
        doc,
        "Interaction diagrams depict message passing between components during runtime execution."
    )

    add_paragraph(
        doc,
        "1. Sequence Diagram - Authentication & JWT Flow: Figure 4.7 traces the user login flow, password verification with bcrypt, "
        "JWT generation, and Bearer token attachment on subsequent protected requests."
    )
    fig_seq_auth = os.path.join(DIAGRAMS_DIR, "fig4_7_sequence_auth.png")
    if os.path.exists(fig_seq_auth):
        add_figure(doc, fig_seq_auth, "Figure 4.7: UML Sequence Diagram - User Authentication & JWT Flow")

    add_paragraph(
        doc,
        "2. Sequence Diagram - PDF Medical Report Analysis Pipeline: Figure 4.8 traces the multipart report upload, in-memory text extraction, "
        "prompt formatting, Groq LLM completion, database logging, and response rendering."
    )
    fig_seq_rep = os.path.join(DIAGRAMS_DIR, "fig4_8_sequence_report.png")
    if os.path.exists(fig_seq_rep):
        add_figure(doc, fig_seq_rep, "Figure 4.8: UML Sequence Diagram - PDF Medical Report Analysis Pipeline")

    add_paragraph(
        doc,
        "3. UML Collaboration / Communication Diagram: Figure 4.9 depicts structural links and numbered message invocations between ClientUI, "
        "ApiRouter, AuthService, PdfService, and DatabaseSession."
    )
    fig_collab = os.path.join(DIAGRAMS_DIR, "fig4_9_collaboration.png")
    if os.path.exists(fig_collab):
        add_figure(doc, fig_collab, "Figure 4.9: UML Collaboration Diagram - Inter-Object Interaction Flow")

    add_heading3(doc, "4.3.4 UML Object Diagram")
    add_paragraph(
        doc,
        "The Object Diagram in Figure 4.10 illustrates runtime instances (User, BearerToken, UploadFile, Activity, GroqSession) and their attribute states."
    )

    fig_obj = os.path.join(DIAGRAMS_DIR, "fig4_10_object_diagram.png")
    if os.path.exists(fig_obj):
        add_figure(doc, fig_obj, "Figure 4.10: UML Object Diagram - Runtime Execution Snapshot")

    add_heading3(doc, "4.3.5 Control Flow / Activity Diagram")
    add_paragraph(
        doc,
        "Figure 4.11 displays the Activity Diagram illustrating the procedural decision flow for request authentication, input validation, "
        "prompt preparation, LLM execution, activity commitment, and response rendering."
    )

    fig_act = os.path.join(DIAGRAMS_DIR, "fig4_11_activity_control_flow.png")
    if os.path.exists(fig_act):
        add_figure(doc, fig_act, "Figure 4.11: Control Flow / Activity Diagram - User Request & AI Guardrail Pipeline")

    add_heading2(doc, "4.4 Database Design & Entity-Relationship (E-R) Diagram")
    add_paragraph(
        doc,
        "The database layer utilizes SQLite 3 managed through SQLAlchemy ORM. The relational model consists of two tables: "
        "users and activities."
    )

    add_heading3(doc, "4.4.1 Entity-Relationship (E-R) Diagram")
    add_paragraph(
        doc,
        "Figure 4.12 displays the Entity-Relationship diagram showing the 1:N relationship between the USERS entity and the ACTIVITIES entity."
    )

    fig_er = os.path.join(DIAGRAMS_DIR, "fig4_12_er_diagram.png")
    if os.path.exists(fig_er):
        add_figure(doc, fig_er, "Figure 4.12: Entity-Relationship (E-R) Diagram - MediMind AI Database Model")

    add_heading3(doc, "4.4.2 Database Tables and Schema Definitions")
    add_paragraph(
        doc,
        "Tables 4.1 and 4.2 define the exact database schemas implemented in the MediMind AI codebase."
    )

    t1_headers = ["Field Name", "Data Type", "Constraints", "Description"]
    t1_rows = [
        ["id", "INTEGER", "PRIMARY KEY, AUTOINCREMENT, INDEX=True", "Unique identifier for each registered user"],
        ["name", "VARCHAR(100)", "NOT NULL", "Full display name of the user"],
        ["email", "VARCHAR(150)", "NOT NULL, UNIQUE=True, INDEX=True", "User email address, used for login lookup"],
        ["password_hash", "VARCHAR(255)", "NOT NULL", "Bcrypt password hash string"]
    ]
    add_table(doc, t1_headers, t1_rows, [1.4, 1.4, 1.8, 2.1])
    add_paragraph(doc, "Table 4.1: Database Schema: USERS Table Definition", italic=True)

    t2_headers = ["Field Name", "Data Type", "Constraints", "Description"]
    t2_rows = [
        ["id", "INTEGER", "PRIMARY KEY, AUTOINCREMENT, INDEX=True", "Unique identifier for each activity entry"],
        ["user_id", "INTEGER", "FOREIGN KEY (users.id), NOT NULL, INDEX=True", "Foreign key reference linking activity to the user"],
        ["feature", "VARCHAR(50)", "NOT NULL", "Feature name (Symptom Checker, Report Analyzer, Medicine Information, AI Health Chat)"],
        ["question", "TEXT", "NULLABLE=True", "User input text or uploaded medical report filename"],
        ["response", "TEXT", "NULLABLE=True", "AI generated response text"],
        ["created_at", "DATETIME", "NOT NULL, DEFAULT=datetime.now(timezone.utc)", "UTC timestamp when the activity was created"]
    ]
    add_table(doc, t2_headers, t2_rows, [1.2, 1.4, 2.0, 2.1])
    add_paragraph(doc, "Table 4.2: Database Schema: ACTIVITIES Table Definition", italic=True)

    add_paragraph(
        doc,
        "Note on Stored Procedures: SQLite is an embedded, serverless relational database engine that does not utilize "
        "database-resident PL/SQL stored procedures. Instead, data validation, transactions, and aggregate computations "
        "are handled cleanly in Python through SQLAlchemy ORM sessions.",
        bold_prefix="Database Stored Procedures Note: "
    )

    doc.add_page_break()