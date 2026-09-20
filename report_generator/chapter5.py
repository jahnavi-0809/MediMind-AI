import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from report_generator.styles import (
    add_heading1, add_heading2, add_heading3, add_paragraph,
    add_bullet, add_table
)

def add_chapter5(doc):
    add_heading1(doc, "CHAPTER 5\nTECHNOLOGY / METHODOLOGY", centered=True)

    add_heading2(doc, "5.1 Architecture & Software Design Patterns")
    add_paragraph(
        doc,
        "The software engineering methodology adopted for MediMind AI integrates industry-standard design patterns across both backend "
        "and frontend ecosystems, guaranteeing maintainability, testability, and high operational throughput:"
    )
    add_bullet(
        doc,
        "FastAPI leverages Python's type hints and the 'Depends' decorator to inject database sessions (get_db) and authenticated user entities "
        "(get_current_user) into router endpoints. This decouples business logic from connection lifecycle management and simplifies mock testing.",
        "Dependency Injection (IoC) Pattern: "
    )
    add_bullet(
        doc,
        "SQLAlchemy 2.0 serves as the data mapping layer, decoupling physical SQL table definitions from domain entity models. Database connection "
        "pooling and transactional session commits/rollbacks are centrally controlled, preventing dangling transactions and connection exhaustion.",
        "Repository & ORM Pattern: "
    )
    add_bullet(
        doc,
        "The React 19 frontend is structured as a tree of modular, reusable functional components (Navbar, Footer, SplashScreen, HealthTips, "
        "TeluguKeyboard). Cross-cutting application concerns (authentication state, JWT tokens, UI language selection) are managed via centralized "
        "React Context Providers (AuthContext, LanguageContext), eliminating prop drilling across nested component hierarchies.",
        "Component-Based Architecture & Context Providers: "
    )
    add_bullet(
        doc,
        "Axios request and response interceptors centrally attach 'Authorization: Bearer <token>' headers to all outgoing API calls. On encountering "
        "an HTTP 401 Unauthorized response, interceptors automatically flush expired tokens from LocalStorage and redirect the client to the login screen.",
        "Interceptor Pattern for HTTP Client: "
    )

    add_heading2(doc, "5.2 System Modules and Their Detailed Functionalities")
    add_paragraph(
        doc,
        "MediMind AI is architectured into seven cohesive, loosely coupled functional modules, each addressing a dedicated clinical or system domain:"
    )

    add_heading3(doc, "Module 1: User Authentication & Session Management Module")
    add_paragraph(
        doc,
        "This module manages user identity lifecycle from registration to session invalidation. Passwords submitted during registration "
        "are sanitized, verified for complexity (minimum 6 characters), and hashed using bcrypt (Blowfish-based salt factor 12). Login requests "
        "validate stored hashes using Passlib CryptContext. Upon successful credential verification, the system issues a cryptographically signed "
        "JSON Web Token (JWT) encoding the user ID as subject ('sub') and an expiration timestamp (60 minutes). The /auth/forgot-password route allows "
        "secure credential updates, while /auth/me returns identity details for the currently authenticated session."
    )

    add_heading3(doc, "Module 2: AI Health Consultation & Conversational Chat Engine")
    add_paragraph(
        doc,
        "The conversational chat module provides real-time, interactive health guidance. Implemented via AIChat.jsx on the frontend and the /ask "
        "endpoint on the backend, it allows patients to submit open-ended questions concerning wellness, symptoms, and disease prevention. The engine "
        "applies intelligent conversational guardrails: responses are strictly capped under 500 words, formatting utilizes bullet points and short paragraphs, "
        "and medical concepts are translated into plain, reassuring language."
    )

    add_heading3(doc, "Module 3: Symptom Checker & Triage Analysis Module")
    add_paragraph(
        doc,
        "The Symptom Checker module (SymptomChecker.jsx) enables patients to describe multi-symptom presentations in natural language. The backend "
        "identifies the 'Symptom Checker' feature via keyword detection ('my symptoms are:') and conditions the Groq LLM to format responses into "
        "standardized clinical triage blocks: Symptoms Summary, Possible Common Causes, General Self-Care, When to Consult a Doctor, and Emergency "
        "Warning Signs. This structure ensures users receive immediate clarity regarding urgency levels."
    )

    add_heading3(doc, "Module 4: In-Memory PDF Medical Report Extraction & Interpretation Module")
    add_paragraph(
        doc,
        "The Medical Report Analyzer (ReportAnalyzer.jsx) provides automated interpretation of complex laboratory documents. Users upload diagnostic "
        "PDFs via a drag-and-drop interface. The backend receives the file as an asynchronous UploadFile, verifies the '%PDF-' magic byte sequence, "
        "and extracts text directly from memory using PyMuPDF (fitz). The extracted report text is embedded into a specialized clinical prompt that "
        "structures findings into: Report Summary, Important Observations, Abnormal Values, General Health Suggestions, Doctor Follow-up Advice, "
        "and Medical Disclaimers."
    )

    add_heading3(doc, "Module 5: Medicine Information & Pharmacovigilance Guidance Module")
    add_paragraph(
        doc,
        "The Medicine Information module (MedicineInfo.jsx) serves as an educational drug encyclopedia. Patients query prescription or over-the-counter "
        "medications. The system outputs structured monographs detailing: Medicine Name, Main Uses, Common Side Effects, Precautions (e.g., pregnancy, "
        "kidney health, alcohol contraindications), and When to Consult a Doctor. Crucially, the module strictly suppresses pharmaceutical dosage "
        "recommendations to prevent accidental overdose or self-medication."
    )

    add_heading3(doc, "Module 6: User Health Activity Tracking & Analytics Module")
    add_paragraph(
        doc,
        "Every clinical query executed across Symptom Checker, Report Analyzer, Medicine Info, and AI Chat is automatically logged as an Activity "
        "record linked to the user's primary key. The /history endpoint returns a reverse-chronological activity log, while the clean_activity_question "
        "utility sanitizes internal prompt prefixes so users see only their original text. The /history/stats endpoint executes optimized SQLAlchemy "
        "aggregate counts (func.count), supplying live analytics (Total Activities, Reports Analyzed, Symptom Checks, AI Chats, Medicine Searches) "
        "directly to the interactive user dashboard."
    )

    add_heading3(doc, "Module 7: Multilingual Engine & Virtual Telugu Keyboard Module")
    add_paragraph(
        doc,
        "This module bridges the vernacular divide through a two-fold localization framework. On the client side, LanguageContext maintains an "
        "internationalization dictionary containing bilingual strings for all UI buttons, headings, placeholders, and error messages. Concurrently, "
        "an interactive on-screen Telugu keyboard component (TeluguKeyboard.jsx) renders character grids containing Telugu vowels (అ, ఆ, ఇ...), "
        "consonants (క, ఖ, గ...), and guninthalu/matras (ా, ి, ీ...), enabling seamless typing without requiring complex OS-level IME configurations."
    )

    add_heading2(doc, "5.3 AI Integration & Prompt Engineering")
    add_paragraph(
        doc,
        "MediMind AI interfaces with Groq Cloud's ultra-low latency Language Processing Unit (LPU) architecture, running the 'openai/gpt-oss-20b' "
        "model. Prompt engineering is leveraged as the primary mechanism for behavioral steering, safety enforcement, and vernacular conditioning."
    )
    add_paragraph(
        doc,
        "The master system prompt in app/services/llm_service.py establishes an immutable clinical persona: 'You are MediMind AI, a professional "
        "AI Healthcare Assistant. Your job is to provide safe, concise, easy-to-understand educational health information.' The prompt specifies "
        "14 non-negotiable operational rules, including: never claiming physician licensure, never generating a final medical diagnosis, never prescribing "
        "pharmaceuticals or dosages, and advising immediate emergency medical intervention for acute symptoms."
    )
    add_paragraph(
        doc,
        "Inference hyper-parameters are fine-tuned for deterministic clinical reasoning: temperature is pinned at 0.2 (minimizing creative drift "
        "and hallucinations) and max_tokens is set to 600 (enforcing concise, focused responses that respect cognitive load)."
    )

    add_heading2(doc, "5.4 In-Memory PDF Processing Pipeline")
    add_paragraph(
        doc,
        "Traditional server-side PDF processing workflows write uploaded multipart files to temporary directories on the disk, invoking external "
        "binaries or libraries that hold operating system file locks. In production environments—especially on Windows servers—file locks cause "
        "frequent 'PermissionError: [WinError 32]' crashes when attempting to clean up temp files concurrently."
    )
    add_paragraph(
        doc,
        "MediMind AI completely redesigns this pipeline in app/services/pdf_service.py by processing PDF files purely in memory. When a multipart file "
        "is uploaded, the asynchronous handler calls 'await file.read()', capturing raw bytes in memory. The system first checks that the byte array "
        "begins with the magic signature b'%PDF-'. If valid, it invokes 'fitz.open(stream=file_bytes, filetype=\"pdf\")'. PyMuPDF iterates across document "
        "pages in RAM, accumulates UTF-8 text strings, and promptly closes the document object. Zero physical files are written to disk, eliminating "
        "disk exhaustion risks, disk I/O latency, and file-lock concurrency faults."
    )

    add_heading2(doc, "5.5 Multilingual Localization & Telugu Virtual Keyboard")
    add_paragraph(
        doc,
        "To ensure authentic vernacular utility rather than awkward literal translation, MediMind AI employs a contextual language conditioning strategy:"
    )
    add_paragraph(
        doc,
        "When a user selects Telugu, the backend dynamically injects language-specific directives into the LLM prompt: 'Respond completely in Telugu "
        "language. Use simple Telugu that elderly users can easily understand. Medical terms can be written in Telugu with the English medical term "
        "in brackets when helpful. Do not translate medical values incorrectly. Keep numbers, measurements, test names, and values accurate.' "
        "This ensures laboratory markers (such as 'Hemoglobin: 13.5 g/dL') remain pristine while the diagnostic interpretation is rendered in natural, "
        "colloquial Telugu."
    )
    add_paragraph(
        doc,
        "The frontend TeluguKeyboard.jsx component complements this by arranging characters into logical orthographic rows. Clicking any virtual key "
        "appends the character to the active React state string, while backspace and clear buttons provide full tactile editing control directly within the browser."
    )

    doc.add_page_break()