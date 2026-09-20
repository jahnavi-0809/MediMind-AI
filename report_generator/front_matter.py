import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from report_generator.styles import (
    PRIMARY_COLOR, DARK_TEXT, add_heading1, add_heading2, add_heading3,
    add_paragraph, add_bullet, add_table, set_cell_background, set_cell_margins
)

def add_cover_page(doc):
    p_spacer1 = doc.add_paragraph()
    p_spacer1.paragraph_format.space_before = Pt(24)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(18)
    r_title = p_title.add_run("MEDIMIND AI: AN INTELLIGENT MULTILINGUAL HEALTHCARE ASSISTANT POWERED BY LLMS AND FASTAPI")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = PRIMARY_COLOR

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(8)
    r_sub = p_sub.add_run("A report submitted in partial fulfillment of the requirements for the Award of Degree of")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True

    p_deg = doc.add_paragraph()
    p_deg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_deg.paragraph_format.space_after = Pt(4)
    r_deg = p_deg.add_run("BACHELOR OF TECHNOLOGY")
    r_deg.font.name = 'Times New Roman'
    r_deg.font.size = Pt(14)
    r_deg.font.bold = True

    p_in = doc.add_paragraph()
    p_in.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_in.paragraph_format.space_after = Pt(4)
    r_in = p_in.add_run("in")
    r_in.font.name = 'Times New Roman'
    r_in.font.size = Pt(12)

    p_dept = doc.add_paragraph()
    p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dept.paragraph_format.space_after = Pt(28)
    r_dept = p_dept.add_run("ARTIFICIAL INTELLIGENCE AND MACHINE LEARNING")
    r_dept.font.name = 'Times New Roman'
    r_dept.font.size = Pt(13)
    r_dept.font.bold = True
    r_dept.font.color.rgb = PRIMARY_COLOR

    p_by = doc.add_paragraph()
    p_by.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_by.paragraph_format.space_after = Pt(6)
    r_by = p_by.add_run("by")
    r_by.font.name = 'Times New Roman'
    r_by.font.size = Pt(12)

    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_after = Pt(4)
    r_name = p_name.add_run("EVURI JAHNAVI")
    r_name.font.name = 'Times New Roman'
    r_name.font.size = Pt(15)
    r_name.font.bold = True
    r_name.font.color.rgb = DARK_TEXT

    p_roll = doc.add_paragraph()
    p_roll.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_roll.paragraph_format.space_after = Pt(40)
    r_roll = p_roll.add_run("Roll Number: Y23ACM422  (Section A)")
    r_roll.font.name = 'Times New Roman'
    r_roll.font.size = Pt(12.5)
    r_roll.font.bold = True

    p_col_dept = doc.add_paragraph()
    p_col_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_col_dept.paragraph_format.space_after = Pt(3)
    r_c1 = p_col_dept.add_run("Department of Cyber Security, Data Science & AIML")
    r_c1.font.name = 'Times New Roman'
    r_c1.font.size = Pt(13)
    r_c1.font.bold = True
    r_c1.font.color.rgb = PRIMARY_COLOR

    p_col = doc.add_paragraph()
    p_col.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_col.paragraph_format.space_after = Pt(3)
    r_c2 = p_col.add_run("Bapatla Engineering College")
    r_c2.font.name = 'Times New Roman'
    r_c2.font.size = Pt(15)
    r_c2.font.bold = True

    p_aut = doc.add_paragraph()
    p_aut.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_aut.paragraph_format.space_after = Pt(3)
    r_c3 = p_aut.add_run("(Autonomous)")
    r_c3.font.name = 'Times New Roman'
    r_c3.font.size = Pt(12)
    r_c3.font.bold = True

    p_aff = doc.add_paragraph()
    p_aff.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_aff.paragraph_format.space_after = Pt(3)
    r_c4 = p_aff.add_run("Affiliated to Acharya Nagarjuna University, Guntur")
    r_c4.font.name = 'Times New Roman'
    r_c4.font.size = Pt(12)
    r_c4.font.bold = True

    p_place = doc.add_paragraph()
    p_place.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_place.paragraph_format.space_after = Pt(3)
    r_c5 = p_place.add_run("Bapatla - 522102, Andhra Pradesh, India")
    r_c5.font.name = 'Times New Roman'
    r_c5.font.size = Pt(12)
    r_c5.font.bold = True

    p_yr = doc.add_paragraph()
    p_yr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_yr.paragraph_format.space_after = Pt(0)
    r_c6 = p_yr.add_run("2026")
    r_c6.font.name = 'Times New Roman'
    r_c6.font.size = Pt(13)
    r_c6.font.bold = True

    doc.add_page_break()

def add_certificate(doc):
    p_hd = doc.add_paragraph()
    p_hd.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_hd.paragraph_format.space_before = Pt(30)
    p_hd.paragraph_format.space_after = Pt(24)
    r_hd = p_hd.add_run("CERTIFICATE")
    r_hd.font.name = 'Times New Roman'
    r_hd.font.size = Pt(16)
    r_hd.font.bold = True
    r_hd.font.color.rgb = PRIMARY_COLOR

    p_body = doc.add_paragraph()
    p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_body.paragraph_format.line_spacing = 1.4
    p_body.paragraph_format.space_after = Pt(40)

    p_body.add_run("This is to certify that the Internship report entitled ")
    r_t = p_body.add_run("“MediMind AI”")
    r_t.bold = True
    p_body.add_run(" being submitted by ")
    r_s = p_body.add_run("EVURI JAHNAVI")
    r_s.bold = True
    p_body.add_run(" (Roll number: ")
    r_r = p_body.add_run("Y23ACM422")
    r_r.bold = True
    p_body.add_run(") is work done by her and submitted during the ")
    r_y = p_body.add_run("2025–2026")
    r_y.bold = True
    p_body.add_run(" academic year, in partial fulfillment of the requirements for the award of the degree of ")
    r_d = p_body.add_run("BACHELOR OF TECHNOLOGY")
    r_d.bold = True
    p_body.add_run(" in ")
    r_b = p_body.add_run("Department of Cyber Security, Data Science & AIML")
    r_b.bold = True
    p_body.add_run(", at ")
    r_co = p_body.add_run("SkillDzire")
    r_co.bold = True
    p_body.add_run(" from ")
    r_d1 = p_body.add_run("25-05-2026")
    r_d1.bold = True
    p_body.add_run(" and ")
    r_d2 = p_body.add_run("06-07-2026")
    r_d2.bold = True
    p_body.add_run(".")

    tbl = doc.add_table(rows=2, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in tbl.rows:
        row.cells[0].width = Inches(2.2)
        row.cells[1].width = Inches(2.2)
        row.cells[2].width = Inches(2.2)

    c0, c1, c2 = tbl.rows[0].cells
    c0.text = "\n\n\n_______________________\nInternal Guide\n[KEEP AS PLACEHOLDER]"
    c1.text = "\n\n\n_______________________\nDepartment Internship\nCoordinator"
    c2.text = "\n\n\n_______________________\nHead of the Department\nCyber Security, DS & AIML"

    for c in [c0, c1, c2]:
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10.5)
            r.font.bold = True

    doc.add_page_break()

def add_declaration(doc):
    p_hd = doc.add_paragraph()
    p_hd.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_hd.paragraph_format.space_before = Pt(36)
    p_hd.paragraph_format.space_after = Pt(30)
    r_hd = p_hd.add_run("DECLARATION")
    r_hd.font.name = 'Times New Roman'
    r_hd.font.size = Pt(16)
    r_hd.font.bold = True
    r_hd.font.color.rgb = PRIMARY_COLOR

    p_body = doc.add_paragraph()
    p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_body.paragraph_format.line_spacing = 1.4
    p_body.paragraph_format.space_after = Pt(60)

    p_body.add_run("I hereby declare that the dissertation entitled ")
    r_t = p_body.add_run("“MediMind AI”")
    r_t.bold = True
    p_body.add_run(" submitted for the award of the degree of ")
    r_d = p_body.add_run("Bachelor of Technology in Artificial Intelligence and Machine Learning")
    r_d.bold = True
    p_body.add_run(" is my original work carried out during my internship at SkillDzire, and the dissertation has not formed the basis for the award of any degree, associates, fellowship, or any other similar titles in any other institution or university.")

    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.rows[0].cells[0].width = Inches(3.3)
    tbl.rows[0].cells[1].width = Inches(3.3)

    c_left = tbl.rows[0].cells[0]
    c_right = tbl.rows[0].cells[1]

    c_left.text = "Place: Bapatla\n\nDate: 06-07-2026"
    c_right.text = "\n\nEVURI JAHNAVI\nReg No: Y23ACM422\nB.Tech (AI & ML), Section A"

    p_l = c_left.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in p_l.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11.5)

    p_r = c_right.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in p_r.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11.5)
        r.font.bold = True

    doc.add_page_break()

def add_company_certificate_placeholder(doc):
    add_heading1(doc, "INTERSHIP CERTIFICATE ISSUED BY COMPANY", centered=True)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=260, bottom=260, left=260, right=260)

    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.line_spacing = 1.4

    r1 = cp.add_run("[INTERNSHIP CERTIFICATE ISSUED BY SKILLDZIRE]\n\n")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(14)
    r1.font.bold = True
    r1.font.color.rgb = PRIMARY_COLOR

    r2 = cp.add_run(
        "[Insert Original Certificate Issued by SkillDzire Here]\n\n"
        "Student Name: EVURI JAHNAVI\n"
        "Roll Number: Y23ACM422 (Section A)\n"
        "Branch: Artificial Intelligence and Machine Learning\n"
        "Internship Organization: SkillDzire\n"
        "Internship Period: 25-05-2026 to 06-07-2026\n"
        "Project Title: MediMind AI\n\n"
        "External Guide: [KEEP AS PLACEHOLDER]\n"
        "Authorized Signatory: [KEEP AS PLACEHOLDER]"
    )
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(11)
    r2.font.italic = True

    doc.add_page_break()

def add_company_profile(doc):
    add_heading1(doc, "COMPANY PROFILE and EXTERNAL GUIDE DETAILS", centered=True)

    add_heading2(doc, "Organization Details")
    headers = ["Parameter", "Details"]
    rows = [
        ["Internship Organization", "SkillDzire"],
        ["Domain", "Skill Development & Technical Internship Program"],
        ["Project Title", "MediMind AI"],
        ["Internship Start Date", "25-05-2026"],
        ["Internship End Date", "06-07-2026"],
        ["Organization Profile", "[Organization profile to be verified / provided by SkillDzire]"]
    ]
    add_table(doc, headers, rows, [2.5, 4.0])

    add_heading2(doc, "External Guide Details")
    g_headers = ["Parameter", "Details"]
    g_rows = [
        ["External Guide Name", "[KEEP AS PLACEHOLDER]"],
        ["Designation", "[KEEP AS PLACEHOLDER]"],
        ["Organization", "SkillDzire"],
        ["Contact / Email", "[KEEP AS PLACEHOLDER]"],
        ["Office Address", "[KEEP AS PLACEHOLDER]"]
    ]
    add_table(doc, g_headers, g_rows, [2.5, 4.0])

    doc.add_page_break()

def add_acknowledgements(doc):
    add_heading1(doc, "ACKNOWLEDGEMENTS", centered=True)

    add_paragraph(
        doc,
        "The completion of this B.Tech internship and the realization of the project 'MediMind AI' has been an "
        "enriching academic and professional journey. I take this opportunity to express my sincere gratitude to all "
        "the individuals and institutions whose guidance and support made this work possible."
    )

    add_paragraph(
        doc,
        "I express my deep gratitude to the Principal of Bapatla Engineering College (Autonomous), for providing "
        "the necessary infrastructural facilities and supportive academic environment throughout my course of study."
    )

    add_paragraph(
        doc,
        "I convey my sincere thanks to the Head of the Department, Department of Cyber Security, Data Science & AIML, "
        "for his constant encouragement and support during the departmental reviews."
    )

    add_paragraph(
        doc,
        "I am grateful to the Department Internship Coordinator, Department of Cyber Security, Data Science & AIML, "
        "for organizing the internship evaluation workflow and coordinating project milestones."
    )

    add_paragraph(
        doc,
        "I extend my heartfelt thanks to my Internal Guide, [KEEP AS PLACEHOLDER], Department of Cyber Security, "
        "Data Science & AIML, for their guidance, scholarly advice, and technical reviews throughout the project."
    )

    add_paragraph(
        doc,
        "I express my gratitude to my External Guide, [KEEP AS PLACEHOLDER], and the team at SkillDzire for providing "
        "the internship opportunity and technical mentorship during the project tenure."
    )

    add_paragraph(
        doc,
        "Finally, I express my gratitude to my parents, family members, and friends for their continuous support "
        "and encouragement throughout my studies."
    )

    p_sig = doc.add_paragraph()
    p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig.paragraph_format.space_before = Pt(24)
    r_sig = p_sig.add_run("EVURI JAHNAVI\n(Roll No: Y23ACM422)\nB.Tech IV Year, AIML - Section A")
    r_sig.font.name = 'Times New Roman'
    r_sig.font.size = Pt(11.5)
    r_sig.font.bold = True

    doc.add_page_break()

def add_abstract(doc):
    add_heading1(doc, "ABSTRACT", centered=True)

    add_paragraph(
        doc,
        "Access to understandable, timely, and culturally accessible healthcare information is a critical challenge in health informatics. "
        "Traditional clinical laboratory reports are filled with technical terminology and numerical indicators that are difficult for laypersons "
        "to interpret, especially in non-English speaking regional communities. Furthermore, general internet health searches often cause unnecessary "
        "anxiety or encourage self-medication without professional guidance. This internship report presents MediMind AI, an intelligent multilingual "
        "healthcare assistant developed during an industry internship at SkillDzire."
    )

    add_paragraph(
        doc,
        "MediMind AI implements a three-tier architecture comprising a FastAPI (Python) asynchronous backend, a responsive React 19 Single Page Application "
        "frontend, a local SQLite relational database managed via SQLAlchemy ORM, and cloud-based Large Language Model inference via the Groq Cloud API "
        "utilizing the 'openai/gpt-oss-20b' model. The system features an in-memory PDF parsing pipeline using PyMuPDF (fitz) that reads uploaded diagnostic "
        "reports directly from memory byte streams, avoiding temporary disk file creation and eliminating Windows file-lock concurrency issues."
    )

    add_paragraph(
        doc,
        "To support regional accessibility, the application provides dual-language capability across English and Telugu, complemented by a custom "
        "virtual Telugu on-screen keyboard. Clinical safety rules are strictly enforced through prompt engineering: the AI organizes responses into "
        "standardized sections (Summary, Important Findings, Abnormal Values, General Suggestions, When to Consult a Doctor, and Disclaimers) while "
        "unconditionally refusing to prescribe medications or dosages and disclaiming medical licensure."
    )

    p_kw = doc.add_paragraph()
    p_kw.paragraph_format.space_before = Pt(14)
    r_kwh = p_kw.add_run("Keywords: ")
    r_kwh.font.name = 'Times New Roman'
    r_kwh.font.size = Pt(11.5)
    r_kwh.font.bold = True
    r_kwh.font.color.rgb = PRIMARY_COLOR

    r_kw = p_kw.add_run(
        "Artificial Intelligence in Healthcare, Large Language Models, FastAPI, React, PyMuPDF, "
        "SQLite, Multilingual Health Assistant, JWT Authentication."
    )
    r_kw.font.name = 'Times New Roman'
    r_kw.font.size = Pt(11)
    r_kw.font.italic = True

    doc.add_page_break()

def add_table_of_contents(doc):
    add_heading1(doc, "TABLE OF CONTENTS", centered=True)

    headers = ["Chapter No.", "Title / Topic", "Page No."]
    rows = [
        ["", "Certificate", "ii"],
        ["", "Declaration", "iii"],
        ["", "Internship Certificate Issued by Company", "iv"],
        ["", "Company Profile and External Guide Details", "v"],
        ["", "Acknowledgements", "vi"],
        ["", "Abstract & Keywords", "vii"],
        ["", "List of Figures", "viii"],
        ["", "List of Tables", "ix"],
        ["1", "INTRODUCTION", "1"],
        ["1.1", "Introduction to Healthcare Informatics & MediMind AI", "1"],
        ["1.2", "Problem Statement", "2"],
        ["1.3", "Objectives of the Project", "3"],
        ["1.4", "Scope of the Project", "3"],
        ["2", "LITERATURE SURVEY / EXISTING SYSTEM", "4"],
        ["2.1", "Existing System Overview", "4"],
        ["2.2", "Limitations of Existing Systems", "4"],
        ["2.3", "Literature Survey", "5"],
        ["2.4", "Proposed System & Key Advantages", "6"],
        ["3", "SOFTWARE REQUIREMENT ANALYSIS", "7"],
        ["3.1", "Software Requirements", "7"],
        ["3.2", "Hardware Requirements", "8"],
        ["3.3", "Functional Requirements Definitions", "8"],
        ["3.4", "Non-Functional Requirements Definitions", "9"],
        ["4", "SOFTWARE DESIGN", "11"],
        ["4.1", "System Architecture", "11"],
        ["4.2", "Data Flow Diagrams (DFD Level 0, Level 1, Level 2)", "12"],
        ["4.3", "UML Diagrams (Class, Sequence, Collaboration, Object, Use Case, Activity)", "15"],
        ["4.4", "Database Design & Entity-Relationship (E-R) Diagram", "20"],
        ["5", "TECHNOLOGY / METHODOLOGY", "22"],
        ["5.1", "Architecture & Design Patterns", "22"],
        ["5.2", "System Modules and Their Detailed Functionalities", "22"],
        ["5.3", "AI Integration & Prompt Engineering", "24"],
        ["5.4", "In-Memory PDF Processing Pipeline", "25"],
        ["5.5", "Multilingual Localization & Telugu Virtual Keyboard", "25"],
        ["6", "CODING / IMPLEMENTATION", "26"],
        ["6.1", "Codebase Organization & File Structure", "26"],
        ["6.2", "Component & Service Functional Walkthrough", "27"],
        ["6.3", "Actual Source Code Listings from MediMind AI", "28"],
        ["6.4", "RESTful API Endpoints Specification", "34"],
        ["6.5", "Authentication Flow & JWT Bearer Token Security", "35"],
        ["7", "TESTING", "36"],
        ["7.1", "Testing Methodologies & Strategy", "36"],
        ["7.2", "Black Box Test Cases", "36"],
        ["7.3", "White Box Test Cases", "38"],
        ["7.4", "AI Safety & Ethical Guardrails Test Cases", "39"],
        ["8", "OUTPUT SCREENS / RESULTS", "40"],
        ["8.1", "User Interfaces and Workflow Walkthrough", "40"],
        ["8.2", "Qualitative Operational & Performance Observations", "42"],
        ["9", "CONCLUSION", "43"],
        ["9.1", "Conclusion", "43"],
        ["9.2", "Limitations of the Current Implementation", "43"],
        ["9.3", "Future Enhancements", "44"],
        ["10", "REFERENCES", "45"],
        ["", "APPENDICES", "46"],
        ["", "Appendix A: Complete REST API Quick-Reference Sheet", "46"],
        ["", "Appendix B: Clinical Safety & AI Disclaimer Protocol", "47"]
    ]
    add_table(doc, headers, rows, [1.1, 4.5, 0.9])
    doc.add_page_break()

def add_list_of_figures(doc):
    add_heading1(doc, "LIST OF FIGURES", centered=True)

    headers = ["Figure No.", "Figure Title / Caption", "Page No."]
    rows = [
        ["Figure 4.1", "MediMind AI - Three-Tier System Architecture", "11"],
        ["Figure 4.2", "Data Flow Diagram (DFD) Level 0 - Context Diagram", "12"],
        ["Figure 4.3", "Data Flow Diagram (DFD) Level 1 - System Flow Diagram", "13"],
        ["Figure 4.4", "Data Flow Diagram (DFD) Level 2 - Medical Report Analyzer (Process 3.0)", "14"],
        ["Figure 4.5", "UML Use Case Diagram - MediMind AI Healthcare Platform", "15"],
        ["Figure 4.6", "UML Class Diagram - Backend Data Models & Core Services", "16"],
        ["Figure 4.7", "UML Sequence Diagram - User Authentication & JWT Flow", "17"],
        ["Figure 4.8", "UML Sequence Diagram - PDF Medical Report Analysis Pipeline", "17"],
        ["Figure 4.9", "UML Collaboration Diagram - Inter-Object Interaction Flow", "18"],
        ["Figure 4.10", "UML Object Diagram - Runtime Execution Snapshot", "19"],
        ["Figure 4.11", "Control Flow / Activity Diagram - User Request & AI Guardrail Pipeline", "20"],
        ["Figure 4.12", "Entity-Relationship (E-R) Diagram - MediMind AI Database Model", "20"]
    ]
    add_table(doc, headers, rows, [1.2, 4.4, 0.9])
    doc.add_page_break()

def add_list_of_tables(doc):
    add_heading1(doc, "LIST OF TABLES", centered=True)

    headers = ["Table No.", "Table Title / Caption", "Page No."]
    rows = [
        ["Table 3.1", "Software Technology Stack & Library Versions", "7"],
        ["Table 3.2", "Hardware System Requirements", "8"],
        ["Table 4.1", "Database Schema: USERS Table Definition", "21"],
        ["Table 4.2", "Database Schema: ACTIVITIES Table Definition", "21"],
        ["Table 6.1", "Complete MediMind AI RESTful API Specification", "34"],
        ["Table 7.1", "Black Box Test Cases: Authentication & Navigation", "36"],
        ["Table 7.2", "Black Box Test Cases: Clinical AI & PDF Processing", "37"],
        ["Table 7.3", "White Box Test Cases: Backend Validation & Service Units", "38"],
        ["Table 7.4", "AI Safety & Ethical Guardrails Verification Test Cases", "39"]
    ]
    add_table(doc, headers, rows, [1.2, 4.4, 0.9])
    doc.add_page_break()