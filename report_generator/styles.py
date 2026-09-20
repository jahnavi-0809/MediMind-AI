import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

PRIMARY_COLOR = RGBColor(14, 111, 120)    # #0E6F78 (MediMind Teal)
DARK_TEXT = RGBColor(28, 26, 22)          # #1C1A16
MUTED_TEXT = RGBColor(100, 116, 139)      # #64748B
BORDER_COLOR = "D8E7E5"
HEADER_BG = "0E6F78"
ALT_ROW_BG = "F8FAFC"

def setup_document():
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.5)   # College standard binding margin
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.27)   # A4
        section.page_height = Inches(11.69)
    return doc

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f"<w:shd {nsdecls('w')} w:fill='{hex_color}'/>")
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_heading1(doc, text, centered=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    if centered:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = PRIMARY_COLOR
    return p

def add_heading2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13.5)
    run.font.bold = True
    run.font.color.rgb = DARK_TEXT
    return p

def add_heading3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = DARK_TEXT
    return p

def add_paragraph(doc, text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.3
    if bold_prefix:
        r_prefix = p.add_run(bold_prefix)
        r_prefix.font.name = 'Times New Roman'
        r_prefix.font.size = Pt(12)
        r_prefix.font.bold = True
        r_prefix.font.color.rgb = DARK_TEXT
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.italic = italic
    run.font.color.rgb = DARK_TEXT
    return p

def add_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    if bold_prefix:
        r_prefix = p.add_run(bold_prefix)
        r_prefix.font.name = 'Times New Roman'
        r_prefix.font.size = Pt(11.5)
        r_prefix.font.bold = True
        r_prefix.font.color.rgb = DARK_TEXT
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11.5)
    run.font.color.rgb = DARK_TEXT
    return p

def add_code_block(doc, code, filename=""):
    if filename:
        p_fn = doc.add_paragraph()
        p_fn.paragraph_format.space_before = Pt(8)
        p_fn.paragraph_format.space_after = Pt(2)
        p_fn.paragraph_format.keep_with_next = True
        r_fn = p_fn.add_run(f"Listing: {filename}")
        r_fn.font.name = 'Consolas'
        r_fn.font.size = Pt(9.5)
        r_fn.font.bold = True
        r_fn.font.color.rgb = PRIMARY_COLOR

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(0)
    cp.paragraph_format.space_after = Pt(0)
    cp.paragraph_format.line_spacing = 1.05
    c_run = cp.add_run(code.strip())
    c_run.font.name = 'Consolas'
    c_run.font.size = Pt(9)
    c_run.font.color.rgb = RGBColor(15, 23, 42)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(6)

def add_figure(doc, image_path, caption):
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(10)
    p_img.paragraph_format.space_after = Pt(4)
    run = p_img.add_run()
    run.add_picture(image_path, width=Inches(5.7))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_after = Pt(12)
    r_cap = p_cap.add_run(caption)
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(10)
    r_cap.font.bold = True
    r_cap.font.italic = True
    r_cap.font.color.rgb = DARK_TEXT

def add_table(doc, headers, rows, col_widths=None):
    tbl = doc.add_table(rows=len(rows)+1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], HEADER_BG)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row_data in enumerate(rows):
        row_cells = tbl.rows[r_idx+1].cells
        bg_color = ALT_ROW_BG if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=140, right=140)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if p.runs:
                p.runs[0].font.name = 'Times New Roman'
                p.runs[0].font.size = Pt(9.5)
                p.runs[0].font.color.rgb = DARK_TEXT

    if col_widths:
        for row in tbl.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(8)
    return tbl